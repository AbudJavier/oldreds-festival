#!/usr/bin/env python3
"""
Festival Page Generator
Genera una pagina HTML de programacion de festival deportivo
a partir de un PDF o Excel + config.json

Uso:
  python generate.py --pdf fixture.pdf --config config.json --output index.html
  python generate.py --xlsx fixture.xlsx --config config.json --output index.html
"""

import argparse
import json
import sys
import re
from pathlib import Path


# ---------------------------------------------------------------------------
# 1. PARSERS — PDF y Excel a lista plana de partidos
# ---------------------------------------------------------------------------

def parse_pdf(pdf_path):
    """
    Parsea un PDF de fixture con formato estandar de hockey:
    Tablas con columnas [Hora, Cat-Zona, EQ1, EQ2] repetidas por cancha.
    Retorna lista de dicts: {hora, categoria, zona, cancha, eq1, eq2}
    """
    import pdfplumber

    matches = []
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables()
            for table in tables:
                if not table or len(table) < 2:
                    continue
                # Detectar tablas de partidos (tienen "Hora" en header)
                header = table[0] if table[0] else []
                header_row = table[1] if len(table) > 1 and table[1] else header
                header_str = ' '.join(str(c) for c in header_row if c)

                if 'Hora' not in header_str:
                    continue

                # Detectar canchas en la primera fila (ej: [None, "CANCHA 1", None, None, None, "CANCHA 2", ...])
                cancha_row = table[0]
                canchas_pos = {}  # col_index -> cancha_number
                if cancha_row:
                    for ci, cell in enumerate(cancha_row):
                        if cell and 'CANCHA' in str(cell).upper():
                            m = re.search(r'CANCHA\s*(\d+)', str(cell), re.IGNORECASE)
                            if m:
                                canchas_pos[ci] = int(m.group(1))

                # Determinar bloques de 4 columnas por cancha
                # Cada bloque: Hora, Cat-Zona, EQ1, EQ2
                block_size = 4
                blocks = []
                cols_per_row = len(header_row) if header_row else 0

                for start in range(0, cols_per_row, block_size):
                    # Encontrar cancha para este bloque
                    cancha = None
                    for ci in range(start, min(start + block_size, cols_per_row)):
                        if ci in canchas_pos:
                            cancha = canchas_pos[ci]
                            break
                    if cancha:
                        blocks.append((start, cancha))

                # Parsear filas de datos (saltar header rows)
                data_start = 2  # Normalmente fila 0=canchas, fila 1=headers
                for row in table[data_start:]:
                    if not row:
                        continue
                    for block_start, cancha in blocks:
                        hora = row[block_start] if block_start < len(row) else None
                        cat_zona = row[block_start + 1] if block_start + 1 < len(row) else None
                        eq1 = row[block_start + 2] if block_start + 2 < len(row) else None
                        eq2 = row[block_start + 3] if block_start + 3 < len(row) else None

                        if not hora or not eq1:
                            continue

                        hora = str(hora).strip()
                        if hora.upper() == 'HORA' or not re.match(r'\d{2}:\d{2}', hora):
                            continue

                        # Detectar RIEGO
                        cat_zona_str = str(cat_zona).strip() if cat_zona else ''
                        if 'RIEGO' in hora.upper() or 'RIEGO' in cat_zona_str.upper():
                            matches.append({
                                'hora': hora,
                                'categoria': 'RIEGO',
                                'zona': '',
                                'cancha': cancha,
                                'eq1': 'RIEGO',
                                'eq2': ''
                            })
                            continue

                        # Parsear categoria y zona
                        categoria = ''
                        zona = ''
                        if cat_zona_str:
                            # Formatos: "S10 - Z A", "S8 - Z E", "SUB 6"
                            parts = cat_zona_str.split(' - ')
                            categoria = parts[0].strip()
                            zona = parts[1].strip() if len(parts) > 1 else ''

                        matches.append({
                            'hora': hora,
                            'categoria': categoria,
                            'zona': zona,
                            'cancha': cancha,
                            'eq1': str(eq1).strip(),
                            'eq2': str(eq2).strip() if eq2 else ''
                        })

    return matches


def parse_excel(xlsx_path):
    """
    Parsea un Excel con formato plano:
    Columnas: Hora | Cancha | Categoria | Zona | Equipo 1 | Equipo 2

    O formato PDF-like con multiples canchas por fila.
    Auto-detecta el formato.
    """
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    matches = []

    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        # Buscar fila header
        header_idx = None
        for i, row in enumerate(rows):
            row_str = ' '.join(str(c).upper() for c in row if c)
            if 'HORA' in row_str and ('EQUIPO' in row_str or 'EQ' in row_str):
                header_idx = i
                break

        if header_idx is None:
            continue

        header = [str(c).strip().upper() if c else '' for c in rows[header_idx]]

        # Detectar formato: plano vs multi-cancha
        hora_count = sum(1 for h in header if h == 'HORA')

        if hora_count <= 1:
            # Formato plano
            col_map = {}
            for ci, h in enumerate(header):
                if 'HORA' in h:
                    col_map['hora'] = ci
                elif 'CANCHA' in h:
                    col_map['cancha'] = ci
                elif 'CAT' in h or 'CATEGORIA' in h:
                    col_map['cat'] = ci
                elif 'ZONA' in h:
                    col_map['zona'] = ci
                elif 'EQ' in h or 'EQUIPO' in h:
                    if 'eq1' not in col_map:
                        col_map['eq1'] = ci
                    else:
                        col_map['eq2'] = ci

            for row in rows[header_idx + 1:]:
                if not row or not row[col_map.get('hora', 0)]:
                    continue
                hora = str(row[col_map['hora']]).strip()
                if not re.match(r'\d{2}:\d{2}', hora):
                    continue

                cancha_val = row[col_map.get('cancha', 0)]
                cancha = int(re.search(r'\d+', str(cancha_val)).group()) if cancha_val and re.search(r'\d+', str(cancha_val)) else 0
                eq1 = str(row[col_map.get('eq1', 0)]).strip() if row[col_map.get('eq1', 0)] else ''
                eq2 = str(row[col_map.get('eq2', 0)]).strip() if row[col_map.get('eq2', 0)] else ''
                cat = str(row[col_map.get('cat', 0)]).strip() if col_map.get('cat') is not None and row[col_map.get('cat', 0)] else ''
                zona = str(row[col_map.get('zona', 0)]).strip() if col_map.get('zona') is not None and row[col_map.get('zona', 0)] else ''

                if 'RIEGO' in eq1.upper() or 'RIEGO' in hora.upper():
                    matches.append({'hora': hora, 'categoria': 'RIEGO', 'zona': '', 'cancha': cancha, 'eq1': 'RIEGO', 'eq2': ''})
                else:
                    # Separar cat-zona si vienen juntos (ej: "S10 - Z A")
                    if ' - ' in cat and not zona:
                        parts = cat.split(' - ')
                        cat = parts[0].strip()
                        zona = parts[1].strip()
                    matches.append({'hora': hora, 'categoria': cat, 'zona': zona, 'cancha': cancha, 'eq1': eq1, 'eq2': eq2})
        else:
            # Formato multi-cancha (similar al PDF)
            # Detectar canchas en fila anterior al header
            cancha_row = rows[header_idx - 1] if header_idx > 0 else None
            canchas_pos = {}
            if cancha_row:
                for ci, cell in enumerate(cancha_row):
                    if cell and 'CANCHA' in str(cell).upper():
                        m = re.search(r'CANCHA\s*(\d+)', str(cell), re.IGNORECASE)
                        if m:
                            canchas_pos[ci] = int(m.group(1))

            block_size = 4
            blocks = []
            for start in range(0, len(header), block_size):
                cancha = None
                for ci in range(start, min(start + block_size, len(header))):
                    if ci in canchas_pos:
                        cancha = canchas_pos[ci]
                        break
                if cancha:
                    blocks.append((start, cancha))

            for row in rows[header_idx + 1:]:
                if not row:
                    continue
                for block_start, cancha in blocks:
                    hora = row[block_start] if block_start < len(row) else None
                    if not hora:
                        continue
                    hora = str(hora).strip()
                    if not re.match(r'\d{2}:\d{2}', hora):
                        continue
                    cat_zona = row[block_start + 1] if block_start + 1 < len(row) else ''
                    eq1 = row[block_start + 2] if block_start + 2 < len(row) else ''
                    eq2 = row[block_start + 3] if block_start + 3 < len(row) else ''
                    if not eq1:
                        continue

                    cat_zona_str = str(cat_zona).strip() if cat_zona else ''
                    categoria, zona = '', ''
                    if ' - ' in cat_zona_str:
                        parts = cat_zona_str.split(' - ')
                        categoria = parts[0].strip()
                        zona = parts[1].strip()
                    else:
                        categoria = cat_zona_str

                    matches.append({
                        'hora': hora,
                        'categoria': categoria,
                        'zona': zona,
                        'cancha': int(cancha),
                        'eq1': str(eq1).strip(),
                        'eq2': str(eq2).strip() if eq2 else ''
                    })

    wb.close()
    return matches


# ---------------------------------------------------------------------------
# 2. ORGANIZADOR — Agrupa partidos por seccion y cancha
# ---------------------------------------------------------------------------

def normalize_cat(cat_str):
    """Normaliza nombres de categoria: S10->SUB 10, S8->SUB 8, etc."""
    cat = cat_str.upper().strip()
    if cat in ('S10', 'SUB10'):
        return 'SUB 10'
    if cat in ('S8', 'SUB8'):
        return 'SUB 8'
    if cat in ('S6', 'SUB6', 'SUB 6'):
        return 'SUB 6'
    return cat


def matches_for_section(matches, section_cfg):
    """Filtra partidos que pertenecen a una seccion del config."""
    canchas = set(section_cfg.get('canchas', []))
    result = []
    for m in matches:
        if m['cancha'] not in canchas:
            continue
        result.append(m)
    return result


def extract_zones(matches, section_cfg):
    """Extrae zonas y equipos para una seccion."""
    canchas = set(section_cfg.get('canchas', []))
    zones = {}
    for m in matches:
        if m['cancha'] not in canchas or m['categoria'] == 'RIEGO':
            continue
        zona = m['zona'] if m['zona'] else normalize_cat(m['categoria'])
        if zona not in zones:
            zones[zona] = set()
        zones[zona].add(m['eq1'])
        if m['eq2']:
            zones[zona].add(m['eq2'])
    # Ordenar equipos dentro de cada zona
    return {z: sorted(list(teams)) for z, teams in sorted(zones.items())}


# ---------------------------------------------------------------------------
# 3. GENERADOR HTML
# ---------------------------------------------------------------------------

def generate_html(matches, config):
    """Genera el HTML completo."""
    cfg = config
    event_date_parts = cfg.get('fecha', '').split()
    # Intentar extraer dia y mes para isEventDay
    day_num = ''
    month_num = ''
    for p in event_date_parts:
        if p.isdigit():
            day_num = p
            break
    month_map = {'enero': 0, 'febrero': 1, 'marzo': 2, 'abril': 3, 'mayo': 4, 'junio': 5,
                 'julio': 6, 'agosto': 7, 'septiembre': 8, 'octubre': 9, 'noviembre': 10, 'diciembre': 11}
    for word in event_date_parts:
        if word.lower() in month_map:
            month_num = str(month_map[word.lower()])
            break

    # Generar secciones HTML
    sections_html = ''
    all_zone_details = []
    for sec in cfg.get('secciones', []):
        sec_matches = matches_for_section(matches, sec)
        zones = extract_zones(sec_matches, sec)

        # Agrupar partidos por cancha
        by_cancha = {}
        for m in sec_matches:
            if m['cancha'] not in by_cancha:
                by_cancha[m['cancha']] = []
            by_cancha[m['cancha']].append(m)

        # Determinar columna zona/cat
        has_mixed_cats = False
        cats_in_section = set()
        for m in sec_matches:
            if m['categoria'] != 'RIEGO':
                cats_in_section.add(normalize_cat(m['categoria']))
        if len(cats_in_section) > 1:
            has_mixed_cats = True

        zona_header = 'Cat.' if has_mixed_cats else 'Zona'

        # Zonas detail
        zone_html = ''
        if zones:
            zone_items = ''
            for zname, teams in zones.items():
                zone_items += '<div class="zona"><b>{}</b>{}</div>\n'.format(zname, ' &middot; '.join(teams))
            zone_html = '<details><summary>Ver zonas {}</summary><div class="zonas">\n{}</div></details>\n'.format(
                sec.get('titulo', ''), zone_items)

        # Tables per cancha
        tables_html = ''
        for cancha in sorted(by_cancha.keys()):
            cancha_matches = by_cancha[cancha]
            rows = ''
            for m in sorted(cancha_matches, key=lambda x: x['hora']):
                if m['eq1'] == 'RIEGO':
                    rows += '<tr class="riego"><td colspan="4">{} &mdash; RIEGO</td></tr>\n'.format(m['hora'])
                    continue
                zona_display = ''
                if has_mixed_cats:
                    nc = normalize_cat(m['categoria'])
                    if nc == 'SUB 8':
                        zona_display = 'S8 ZE' if m['zona'] == 'Z E' else 'S8'
                    elif nc == 'SUB 6':
                        zona_display = 'SUB 6'
                    else:
                        zona_display = m['zona']
                else:
                    zona_display = m['zona']
                rows += '<tr><td class="h">{}</td><td class="z">{}</td><td>{}</td><td>{}</td></tr>\n'.format(
                    m['hora'], zona_display, m['eq1'], m['eq2'])

            tables_html += '<h3>Cancha {}</h3>\n'.format(cancha)
            tables_html += '<table><tr><th>Hora</th><th>{}</th><th>Equipo 1</th><th>Equipo 2</th></tr>\n{}</table>\n\n'.format(
                zona_header, rows)

        sections_html += '''<section id="{id}">
<h2>{titulo}</h2>
<p class="cat-note">{nota} &mdash; elige tu equipo arriba para destacar tus partidos</p>
{zones}
{tables}
</section>

'''.format(
            id=sec['id'],
            titulo=sec['titulo'],
            nota=sec.get('nota', ''),
            zones=zone_html,
            tables=tables_html
        )

    # Generar mapa
    mapa_cfg = cfg.get('mapa', {})
    mapa_grid = ''
    for fila in mapa_cfg.get('filas', []):
        for cell in fila:
            span = ' style="grid-column:span {}"'.format(cell['span']) if cell.get('span') else ''
            color_class = 'lb' if cell.get('color') == 'b' else 'lg'
            mapa_grid += '<div class="lc {cls}"{span}>CANCHA {n}</div>'.format(
                cls=color_class, span=span, n=cell['cancha'])
        mapa_grid += '\n'

    # Generar colores de clubes para JS
    clubes_js = json.dumps(cfg.get('clubes', {}))

    # Secciones tinte CSS
    tinte_css = ''
    for sec in cfg.get('secciones', []):
        tinte_css += '#{id}{{background:{tinte};border-top:3px solid {borde}}}\n'.format(
            id=sec['id'], tinte=sec.get('tinte', '#fff'), borde=sec.get('borde', '#eee'))

    # Logo
    logo_src = cfg.get('logo_base64', '')
    logo_img = ''
    if logo_src:
        if logo_src.startswith('data:'):
            logo_img = '<img src="{}" alt="{}">'.format(logo_src, cfg.get('organizador', ''))
        else:
            logo_img = '<img src="data:image/png;base64,{}" alt="{}">'.format(logo_src, cfg.get('organizador', ''))

    # isEventDay
    event_day_check = 'return true;'  # default: siempre activo
    if day_num and month_num:
        event_day_check = 'var d=new Date();return d.getMonth()==={m}&&(d.getDate()==={d}||d.getDate()==={dp});'.format(
            m=month_num, d=day_num, dp=int(day_num) - 1)

    color = cfg.get('color_principal', '#CD5241')
    dark = cfg.get('color_oscuro', '#141414')

    html = '''<!DOCTYPE html>
<html lang="es">
<head>
<meta charset="UTF-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{evento} &middot; {org} &middot; {fecha}</title>
<meta property="og:title" content="{evento} {org} &mdash; {fecha}">
<meta property="og:description" content="Programacion completa. Busca tu equipo y encuentra tus horarios y canchas. {lugar}.">
<meta property="og:type" content="website">
<meta property="og:url" content="{url}">
<meta name="theme-color" content="{color}">
<style>
:root{{--red:{color};--dark:{dark};--ink:#2a2a2a;--paper:#faf8f6;--line:#e6e0dc;--riego:#dff0d8}}
*{{margin:0;padding:0;box-sizing:border-box}}
body{{font-family:-apple-system,'Segoe UI',Roboto,sans-serif;background:var(--paper);color:var(--ink);padding-bottom:60px}}
header{{background:var(--dark);color:#fff;padding:22px 16px 18px;text-align:center;border-bottom:4px solid var(--red);position:relative;overflow:hidden}}
header::before{{content:"{org_upper}";position:absolute;top:50%;left:50%;transform:translate(-50%,-50%);font-size:min(64px,7vw);font-weight:900;letter-spacing:8px;text-transform:uppercase;color:rgba(255,255,255,.1);white-space:nowrap;pointer-events:none}}
header img{{width:74px;display:block;margin:0 auto 8px}}
header h1{{font-size:20px;letter-spacing:.5px;text-transform:uppercase}}
header p{{font-size:13px;color:#d8d0cc;margin-top:4px}}
section{{padding:18px 12px 6px;max-width:640px;margin:0 auto}}
h2{{font-size:16px;text-transform:uppercase;letter-spacing:1px;color:var(--dark);border-left:4px solid var(--red);padding-left:10px;margin-bottom:4px}}
.cat-note{{font-size:12px;color:#8a817c;margin:0 0 12px 14px}}
h3{{font-size:13px;text-transform:uppercase;letter-spacing:.5px;background:var(--dark);color:#fff;padding:8px 12px;border-radius:8px 8px 0 0;margin-top:16px}}
table{{width:100%;border-collapse:collapse;background:#fff;font-size:13px;border:1px solid var(--line);border-top:none;border-radius:0 0 8px 8px;overflow:hidden}}
th{{font-size:11px;text-transform:uppercase;color:#8a817c;text-align:left;padding:6px 8px;border-bottom:1px solid var(--line)}}
td{{padding:7px 8px;border-bottom:1px solid #f1ecea}}
tr:last-child td{{border-bottom:none}}
td.h{{font-variant-numeric:tabular-nums;font-weight:600;width:48px}}
td.z{{color:#8a817c;font-size:11px;width:52px}}
tr.oreds{{background:#fdf0ee;box-shadow:inset 3px 0 0 var(--red)}}
tr.oreds td{{font-weight:600}}
tr.riego td{{background:var(--riego);text-align:center;font-weight:700;font-size:11px;letter-spacing:2px;color:#3c6e3c}}
details{{margin:10px 0;background:#fff;border:1px solid var(--line);border-radius:8px}}
summary{{padding:10px 12px;font-size:13px;font-weight:600;cursor:pointer;color:var(--dark)}}
.zonas{{display:flex;flex-wrap:wrap;gap:8px;padding:0 12px 12px}}
.zona{{flex:1 1 130px;font-size:12px;background:var(--paper);border:1px solid var(--line);border-radius:6px;padding:8px}}
.zona b{{display:block;color:var(--red);margin-bottom:4px;font-size:11px;text-transform:uppercase}}
.lgrid{{display:grid;grid-template-columns:1fr 1fr 1fr 1.4fr;gap:6px;font-size:11px;font-weight:700;text-align:center}}
.lc{{border-radius:6px;padding:14px 4px;color:#fff}}
.lb{{background:#5b8db8}}.lg{{background:#7aa874}}
.lnote{{font-size:11px;color:#8a817c;margin-top:8px;text-align:center}}
footer{{text-align:center;font-size:11px;color:#a89f9a;padding:24px 0 10px}}
.finder{{max-width:640px;margin:14px auto 0;padding:0 12px}}
.finder-card{{background:var(--dark);border-radius:12px;padding:16px;color:#fff}}
.finder-card label{{display:block;font-size:11px;text-transform:uppercase;letter-spacing:1px;color:#d8d0cc;margin:10px 0 4px}}
.finder-card label:first-child{{margin-top:0}}
.finder-card select{{width:100%;padding:11px 12px;font-size:15px;border-radius:8px;border:none;background:#fff;color:var(--ink)}}
.fc-title{{font-size:15px;font-weight:700;text-transform:uppercase;letter-spacing:.5px}}
.fc-sub{{font-size:12px;color:#d8d0cc;margin:2px 0 4px}}
.mis{{background:#fff;border:2px solid var(--red);border-radius:12px;margin-top:12px;overflow:hidden;display:none}}
.mis.on{{display:block}}
.mis-head{{background:var(--red);color:#fff;padding:10px 14px;font-weight:700;font-size:14px;text-transform:uppercase;letter-spacing:.5px}}
.mis-item{{display:flex;align-items:center;gap:10px;padding:10px 14px;border-bottom:1px solid #f1ecea;font-size:14px}}
.mis-item:last-child{{border-bottom:none}}
.mis-hora{{font-weight:700;font-variant-numeric:tabular-nums;width:50px}}
.mis-rival{{flex:1}}
.mis-rival small{{display:block;color:#8a817c;font-size:11px}}
.badge{{font-size:11px;font-weight:700;color:#fff;padding:5px 9px;border-radius:6px;white-space:nowrap}}
.badge.g{{background:#7aa874}}.badge.b{{background:#5b8db8}}
.mis-map{{padding:10px 14px;font-size:12px;background:#fdf0ee;color:var(--ink)}}
tr.mine{{background:#fdf0ee;box-shadow:inset 3px 0 0 var(--red)}}
tr.mine td{{font-weight:700}}
.lc{{position:relative;transition:all .2s}}
.lc.mine{{outline:4px solid var(--red);outline-offset:-2px;font-size:13px}}
.lc.mine::after{{content:"TU CANCHA";position:absolute;top:-9px;left:50%;transform:translateX(-50%);background:var(--red);color:#fff;font-size:8px;letter-spacing:1px;padding:2px 6px;border-radius:4px}}
.lc.dim{{opacity:.35}}
.badge-now{{background:#e74c3c;color:#fff;font-size:10px;font-weight:800;padding:3px 7px;border-radius:4px;flex-shrink:0;animation:pulse 1.5s infinite}}
.badge-next{{background:#f39c12;color:#fff;font-size:10px;font-weight:800;padding:3px 7px;border-radius:4px;flex-shrink:0}}
@keyframes pulse{{0%,100%{{opacity:1}}50%{{opacity:.5}}}}
.share-btn{{display:none;width:100%;margin-top:10px;padding:12px;background:var(--red);color:#fff;border:none;border-radius:8px;font-size:14px;font-weight:700;cursor:pointer;text-transform:uppercase;letter-spacing:.5px}}
.share-btn:active{{opacity:.8}}
.team-dot{{display:inline-block;width:10px;height:10px;border-radius:50%;margin-right:6px;vertical-align:middle}}
.mis-sep{{padding:10px 14px;font-size:14px;font-weight:700;text-transform:uppercase;letter-spacing:.5px;color:var(--ink);background:var(--paper);border-bottom:1px solid var(--line)}}
.mis-team-label{{font-size:12px;color:#8a817c;margin-left:8px;font-weight:600;text-transform:none;letter-spacing:0}}
.mis-cat-sep{{background:var(--dark);color:#fff;font-size:13px;letter-spacing:1px;border-bottom:none}}
.mis-group{{border:2px solid var(--line);border-radius:10px;margin:10px 10px 0;overflow:hidden}}
.mis-group .mis-sep{{border-radius:0}}
.mis-group .mis-item:last-child{{border-bottom:none}}
.mis-group-head{{padding:8px 12px;font-size:13px;font-weight:700;text-transform:uppercase;letter-spacing:1px;color:#fff}}
.mis-cancha-head{{padding:6px 12px;font-size:12px;font-weight:700;color:#fff;text-transform:uppercase;letter-spacing:.5px}}
{tinte_css}
.finder-map{{margin-top:12px;background:#fff;border:1px solid var(--line);border-radius:8px;overflow:hidden}}
.finder-map summary{{padding:10px 12px;font-size:13px;font-weight:700;cursor:pointer;color:var(--dark);text-transform:uppercase;letter-spacing:.5px}}
.finder-map .lgrid{{padding:12px}}
.finder-map .lnote{{padding:0 12px 12px}}
</style>
</head>
<body>
<header>
{logo}
<h1>{evento}</h1>
<p>{cats_display} &mdash; {fecha} &middot; {lugar} &middot; {org}</p>
</header>

<div class="finder">
<div class="finder-card">
<div class="fc-title">Encuentra a tu equipo</div>
<p class="fc-sub">Elige categoria y equipo: veras tus horarios y tu cancha marcada en el mapa.</p>
<label for="selCat">Categoria</label>
<select id="selCat"><option value="">— Elegir categoria —</option></select>
<label for="selTeam">Equipo</label>
<select id="selTeam" disabled><option value="">— Primero elige categoria —</option></select>
</div>
<div class="mis" id="misCard">
<div class="mis-head" id="misHead"></div>
<div id="misList"></div>
<div class="mis-map">&#x1F4CD; Tus canchas estan destacadas en el mapa de abajo.</div>
<button class="share-btn" id="shareBtn">&#x1F4E4; Compartir programacion</button>
</div>
<details class="finder-map">
<summary>&#x1F4CD; Mapa de canchas</summary>
<div class="lgrid" id="layout">
{mapa_grid}
</div>
<p class="lnote">{mapa_nota}</p>
</details>
</div>

{sections}

<footer>{org} &middot; {evento} &middot; {lugar}</footer>
<script>
(function(){{
  var TC={clubes_js};
  function tColor(n){{var k=Object.keys(TC);for(var i=0;i<k.length;i++){{if(n.indexOf(k[i])===0)return TC[k[i]];}}return '#bbb';}}
  function clubOf(name){{var m=name.match(/^(.+?)\\s*\\d*$/);return m?m[1].trim():name;}}
  function allClubs(){{var s={{}};MATCHES.forEach(function(x){{s[clubOf(x.eq1)]=1;s[clubOf(x.eq2)]=1;}});return Object.keys(s).sort();}}

  var MATCHES=[];
  document.querySelectorAll('h3').forEach(function(h){{
    var m=h.textContent.match(/Cancha\\s+(\\d)/i);
    if(!m) return;
    var cancha=m[1];
    var table=h.nextElementSibling;
    if(!table||table.tagName!=='TABLE') return;
    var sec=h.closest('section');
    table.querySelectorAll('tr').forEach(function(tr){{
      var td=tr.querySelectorAll('td');
      if(td.length<4) return;
      var cat;
      {cat_detection_js}
      MATCHES.push({{cat:cat,hora:td[0].textContent.trim(),eq1:td[2].textContent.trim(),eq2:td[3].textContent.trim(),cancha:cancha,row:tr}});
    }});
  }});

  var CATS={cats_js};
  var selCat=document.getElementById('selCat'),selTeam=document.getElementById('selTeam');
  var misCard=document.getElementById('misCard'),misHead=document.getElementById('misHead'),misList=document.getElementById('misList');
  var shareBtn=document.getElementById('shareBtn');
  var finderMap=document.querySelector('.finder-map');
  CATS.forEach(function(c){{var o=document.createElement('option');o.value=c;o.textContent=c;selCat.appendChild(o);}});
  var oc=document.createElement('option');oc.value='CLUB';oc.textContent='POR CLUB';selCat.appendChild(oc);
  var oh=document.createElement('option');oh.value='HORARIO';oh.textContent='POR HORARIO';selCat.appendChild(oh);

  function teamsFor(cat){{var s={{}};MATCHES.forEach(function(x){{if(x.cat===cat){{s[x.eq1]=1;s[x.eq2]=1;}}}});return Object.keys(s).sort();}}
  selCat.addEventListener('change',function(){{
    selTeam.innerHTML='<option value="">— Elegir equipo —</option>';
    if(!selCat.value){{selTeam.disabled=true;clearAll();return;}}
    if(selCat.value==='HORARIO'){{
      selTeam.innerHTML='<option value="">— Elegir bloque —</option>';
      var hrs={{}};MATCHES.forEach(function(x){{hrs[x.hora]=1;}});
      Object.keys(hrs).sort().forEach(function(h){{var o=document.createElement('option');o.value=h;o.textContent=h;selTeam.appendChild(o);}});
      selTeam.disabled=false;clearAll();return;
    }}
    var items=(selCat.value==='CLUB')?allClubs():teamsFor(selCat.value);
    items.forEach(function(t){{var o=document.createElement('option');o.value=t;o.textContent=t;selTeam.appendChild(o);}});
    selTeam.disabled=false;clearAll();
  }});
  selTeam.addEventListener('change',function(){{
    if(selCat.value==='CLUB') applyClub(selTeam.value);
    else if(selCat.value==='HORARIO') applyHora(selTeam.value);
    else apply(selCat.value,selTeam.value);
    save();
  }});

  function mapCells(){{return document.querySelectorAll('.lgrid .lc');}}
  function clearAll(){{
    document.querySelectorAll('tr.mine').forEach(function(r){{r.classList.remove('mine');}});
    mapCells().forEach(function(c){{c.classList.remove('mine','dim');}});
    misCard.classList.remove('on');misList.innerHTML='';misHead.textContent='';misHead.style.background='';
    if(shareBtn){{shareBtn.style.display='none';shareBtn.textContent='\\uD83D\\uDCE4 Compartir programacion';}}
  }}
  function openMap(){{if(finderMap&&!finderMap.open) finderMap.open=true;}}
  function badgeClass(n){{return (n==='1'||n==='2'||n==='3')?'g':'b';}}
  function toMins(h){{var p=h.split(':');return parseInt(p[0])*60+parseInt(p[1]);}}
  function isEventDay(){{{event_day_js}}}

  function apply(cat,team){{
    clearAll();if(!cat||!team) return;
    var mine=MATCHES.filter(function(x){{return x.cat===cat&&(x.eq1===team||x.eq2===team);}});
    mine.sort(function(a,b){{return a.hora<b.hora?-1:1;}});
    if(!mine.length) return;
    var live=isEventDay(),nowM=0,ahoraIdx=-1,sigIdx=-1;
    if(live){{nowM=new Date().getHours()*60+new Date().getMinutes();for(var i=0;i<mine.length;i++){{var mm=toMins(mine[i].hora);var end=(i<mine.length-1)?toMins(mine[i+1].hora):mm+15;if(nowM>=mm&&nowM<end)ahoraIdx=i;if(sigIdx===-1&&mm>nowM)sigIdx=i;}}}}
    misHead.textContent=team+' \\u00B7 '+cat+' \\u2014 '+mine.length+' partidos';misHead.style.background=tColor(team);
    misList.innerHTML='';var canchas={{}};
    mine.forEach(function(x,i){{canchas[x.cancha]=1;x.row.classList.add('mine');var rival=(x.eq1===team)?x.eq2:x.eq1;var timeBadge='';if(i===ahoraIdx)timeBadge='<span class="badge-now">AHORA</span>';else if(i===sigIdx)timeBadge='<span class="badge-next">SIGUE</span>';var dot='<span class="team-dot" style="background:'+tColor(rival)+'"></span>';var d=document.createElement('div');d.className='mis-item';d.innerHTML='<span class="mis-hora">'+x.hora+'</span>'+timeBadge+'<span class="mis-rival">'+dot+'vs '+rival+'</span>'+'<span class="badge '+badgeClass(x.cancha)+'">CANCHA '+x.cancha+'</span>';misList.appendChild(d);}});
    mapCells().forEach(function(c){{var mm=c.textContent.match(/CANCHA\\s+(\\d)/);if(mm&&canchas[mm[1]])c.classList.add('mine');else c.classList.add('dim');}});
    misCard.classList.add('on');openMap();if(shareBtn)shareBtn.style.display='block';
  }}

  function applyClub(club){{
    clearAll();if(!club) return;
    var color=tColor(club);
    var mine=MATCHES.filter(function(x){{return clubOf(x.eq1)===club||clubOf(x.eq2)===club;}});
    if(!mine.length) return;
    var teams={{}},teamOrder=[];var catOrder={cat_order_js};
    mine.forEach(function(x){{var t=(clubOf(x.eq1)===club)?x.eq1:x.eq2;var key=x.cat+'|'+t;if(!teams[key]){{teams[key]={{name:t,cat:x.cat,matches:[]}};teamOrder.push(key);}}teams[key].matches.push(x);}});
    teamOrder.sort(function(a,b){{var ga=teams[a],gb=teams[b];var ca=(catOrder[ga.cat]||9),cb=(catOrder[gb.cat]||9);if(ca!==cb)return ca-cb;return ga.name<gb.name?-1:1;}});
    misHead.textContent=club+' \\u2014 '+teamOrder.length+' equipos \\u00B7 '+mine.length+' partidos';misHead.style.background=color;
    misList.innerHTML='';var canchas={{}};var live=isEventDay(),nowM=0;if(live)nowM=new Date().getHours()*60+new Date().getMinutes();var lastCat='',catBox=null;
    teamOrder.forEach(function(tk){{var grp=teams[tk];if(grp.cat!==lastCat){{lastCat=grp.cat;catBox=document.createElement('div');catBox.className='mis-group';catBox.style.borderColor=color;var head=document.createElement('div');head.className='mis-group-head';head.style.background=color;head.textContent=grp.cat;catBox.appendChild(head);misList.appendChild(catBox);}}var sep=document.createElement('div');sep.className='mis-sep';sep.innerHTML=grp.name;catBox.appendChild(sep);var sorted=grp.matches.sort(function(a,b){{return a.hora<b.hora?-1:1;}});var ahoraIdx=-1,sigIdx=-1;if(live){{for(var i=0;i<sorted.length;i++){{var mm=toMins(sorted[i].hora);var end=(i<sorted.length-1)?toMins(sorted[i+1].hora):mm+15;if(nowM>=mm&&nowM<end)ahoraIdx=i;if(sigIdx===-1&&mm>nowM)sigIdx=i;}}}}sorted.forEach(function(x,i){{canchas[x.cancha]=1;x.row.classList.add('mine');var rival=(clubOf(x.eq1)===club)?x.eq2:x.eq1;var timeBadge='';if(i===ahoraIdx)timeBadge='<span class="badge-now">AHORA</span>';else if(i===sigIdx)timeBadge='<span class="badge-next">SIGUE</span>';var dot='<span class="team-dot" style="background:'+tColor(rival)+'"></span>';var d=document.createElement('div');d.className='mis-item';d.innerHTML='<span class="mis-hora">'+x.hora+'</span>'+timeBadge+'<span class="mis-rival">'+dot+'vs '+rival+'</span>'+'<span class="badge '+badgeClass(x.cancha)+'">CANCHA '+x.cancha+'</span>';catBox.appendChild(d);}});}});
    mapCells().forEach(function(c){{var mm=c.textContent.match(/CANCHA\\s+(\\d)/);if(mm&&canchas[mm[1]])c.classList.add('mine');else c.classList.add('dim');}});
    misCard.classList.add('on');openMap();if(shareBtn)shareBtn.style.display='block';
  }}

  function applyHora(hora){{
    clearAll();if(!hora) return;
    var mine=MATCHES.filter(function(x){{return x.hora===hora;}});if(!mine.length) return;
    mine.sort(function(a,b){{return a.cancha<b.cancha?-1:1;}});
    var byCourt={{}},courtOrder=[];
    mine.forEach(function(x){{if(!byCourt[x.cancha]){{byCourt[x.cancha]=[];courtOrder.push(x.cancha);}}byCourt[x.cancha].push(x);}});
    misHead.textContent=hora+' \\u2014 '+mine.length+' partidos en '+courtOrder.length+' canchas';misHead.style.background='{dark}';
    misList.innerHTML='';var canchas={{}};
    courtOrder.forEach(function(cn){{var box=document.createElement('div');box.className='mis-group';var bc=badgeClass(cn);var bgColor=(bc==='g')?'#7aa874':'#5b8db8';box.style.borderColor=bgColor;var head=document.createElement('div');head.className='mis-cancha-head';head.style.background=bgColor;head.textContent='CANCHA '+cn;box.appendChild(head);canchas[cn]=1;byCourt[cn].forEach(function(x){{x.row.classList.add('mine');var dot1='<span class="team-dot" style="background:'+tColor(x.eq1)+'"></span>';var dot2='<span class="team-dot" style="background:'+tColor(x.eq2)+'"></span>';var d=document.createElement('div');d.className='mis-item';d.innerHTML='<span class="mis-rival">'+dot1+x.eq1+'</span>'+'<span style="color:#8a817c;font-size:12px;padding:0 6px">vs</span>'+'<span class="mis-rival">'+dot2+x.eq2+'</span>'+'<span class="badge" style="background:#8a817c;font-size:10px">'+x.cat+'</span>';box.appendChild(d);}});misList.appendChild(box);}});
    mapCells().forEach(function(c){{var mm=c.textContent.match(/CANCHA\\s+(\\d)/);if(mm&&canchas[mm[1]])c.classList.add('mine');else c.classList.add('dim');}});
    misCard.classList.add('on');openMap();if(shareBtn)shareBtn.style.display='block';
  }}

  if(shareBtn) shareBtn.addEventListener('click',function(){{
    var shareText='Programacion '+(selCat.value==='CLUB'?selTeam.value:selTeam.value+' ('+selCat.value+')')+' \\u2014 {evento} {org}, {fecha}, {lugar}';
    var shareUrl='{url}';
    if(navigator.share){{navigator.share({{title:selTeam.value+' \\u2014 {evento} {org}',text:shareText,url:shareUrl}}).catch(function(){{}});}}
    else{{navigator.clipboard.writeText(shareText+' '+shareUrl).then(function(){{shareBtn.textContent='\\u2705 Enlace copiado';setTimeout(function(){{shareBtn.textContent='\\uD83D\\uDCE4 Compartir programacion';}},2000);}}).catch(function(){{}});}}
  }});

  function save(){{try{{localStorage.setItem('festival_team',JSON.stringify({{c:selCat.value,t:selTeam.value}}));}}catch(e){{}}}}
  try{{var prev=JSON.parse(localStorage.getItem('festival_team')||'null');if(prev&&prev.c&&prev.t){{selCat.value=prev.c;selCat.dispatchEvent(new Event('change'));selTeam.value=prev.t;if(selTeam.value===prev.t){{if(prev.c==='CLUB')applyClub(prev.t);else if(prev.c==='HORARIO')applyHora(prev.t);else apply(prev.c,prev.t);}}}}}}catch(e){{}}
  setInterval(function(){{if(selCat.value&&selTeam.value){{if(selCat.value==='CLUB')applyClub(selTeam.value);else if(selCat.value==='HORARIO')applyHora(selTeam.value);else apply(selCat.value,selTeam.value);}}}},30000);
}})();
</script>
</body>
</html>'''

    # Build category detection JS
    secciones = cfg.get('secciones', [])
    cat_lines = []
    for i, sec in enumerate(secciones):
        condition = "sec.id==='{}'".format(sec['id'])
        cat_label = sec['titulo'].split()[0] + ' ' + sec['titulo'].split()[1] if len(sec['titulo'].split()) > 1 else sec['titulo']
        # Use simple labels
        cat_labels = {
            'sub10': 'SUB 10',
            'sub8': 'SUB 8',
            'sub6': 'SUB 6'
        }
        label = cat_labels.get(sec['id'], sec['titulo'])
        if i == 0:
            cat_lines.append("if({}) cat='{}';".format(condition, label))
        elif i < len(secciones) - 1:
            cat_lines.append("else if({}) cat='{}';".format(condition, label))
        else:
            # Last section: check for mixed categories
            cat_lines.append("else cat=(td[1].textContent.trim().indexOf('SUB 6')===0)?'SUB 6':'SUB 8';")
    cat_detection = '\n      '.join(cat_lines)

    # Cats list
    cats_list = []
    for sec in secciones:
        cat_labels = {'sub10': 'SUB 10', 'sub8': 'SUB 8', 'sub6': 'SUB 6'}
        label = cat_labels.get(sec['id'], sec['titulo'])
        if label not in cats_list:
            cats_list.append(label)

    # Cat order for club view
    cat_order = {}
    for i, c in enumerate(cats_list):
        cat_order[c] = i + 1

    # Categories display string
    cats_display = ' &middot; '.join(cats_list)

    return html.format(
        evento=cfg.get('evento', 'Festival'),
        org=cfg.get('organizador', ''),
        org_upper=cfg.get('organizador', '').upper(),
        fecha=cfg.get('fecha', ''),
        lugar=cfg.get('lugar', ''),
        url=cfg.get('url', ''),
        color=color,
        dark=dark,
        logo=logo_img,
        cats_display=cats_display,
        tinte_css=tinte_css,
        mapa_grid=mapa_grid,
        mapa_nota=mapa_cfg.get('nota', ''),
        sections=sections_html,
        clubes_js=clubes_js,
        cat_detection_js=cat_detection,
        cats_js=json.dumps(cats_list),
        cat_order_js=json.dumps(cat_order),
        event_day_js=event_day_check,
    )


# ---------------------------------------------------------------------------
# 4. MAIN
# ---------------------------------------------------------------------------

def main():
    parser = argparse.ArgumentParser(description='Genera pagina HTML de festival deportivo')
    group = parser.add_mutually_exclusive_group(required=True)
    group.add_argument('--pdf', help='Archivo PDF con el fixture')
    group.add_argument('--xlsx', help='Archivo Excel con el fixture')
    parser.add_argument('--config', required=True, help='Archivo config.json')
    parser.add_argument('--output', default='index.html', help='Archivo de salida (default: index.html)')
    args = parser.parse_args()

    # Leer config
    with open(args.config) as f:
        config = json.load(f)

    # Parsear datos
    if args.pdf:
        print('Leyendo PDF: {}'.format(args.pdf))
        matches = parse_pdf(args.pdf)
    else:
        print('Leyendo Excel: {}'.format(args.xlsx))
        matches = parse_excel(args.xlsx)

    print('Partidos encontrados: {}'.format(len(matches)))
    riego = sum(1 for m in matches if m['categoria'] == 'RIEGO')
    real = len(matches) - riego
    print('  Partidos reales: {}, Riegos: {}'.format(real, riego))

    # Equipos unicos
    teams = set()
    for m in matches:
        if m['categoria'] != 'RIEGO':
            teams.add(m['eq1'])
            if m['eq2']:
                teams.add(m['eq2'])
    print('  Equipos unicos: {}'.format(len(teams)))

    # Canchas
    canchas = set(m['cancha'] for m in matches)
    print('  Canchas: {}'.format(sorted(canchas)))

    # Generar HTML
    html = generate_html(matches, config)
    with open(args.output, 'w', encoding='utf-8') as f:
        f.write(html)
    print('\nGenerado: {} ({:.1f} KB)'.format(args.output, len(html) / 1024))


if __name__ == '__main__':
    main()
